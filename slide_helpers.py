"""
Shared helpers and color palette for AI-Native Software Development deck.
All parts import this module so diagrams, fonts, and colors stay consistent.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
import copy, os

# ── Slide dimensions (widescreen 13.33 x 7.5 in) ──────────────────────
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ── Academic colour palette (NO dark colours) ──────────────────────────
BG_COLOR       = RGBColor(0xE7, 0xF2, 0xFA)   # very light blue bg
TITLE_BAR      = RGBColor(0x2F, 0x75, 0xB5)   # medium blue
ACCENT_BLUE    = RGBColor(0x44, 0x72, 0xC4)
ACCENT_GREEN   = RGBColor(0x70, 0xAD, 0x47)
ACCENT_ORANGE  = RGBColor(0xED, 0x7D, 0x31)
ACCENT_PURPLE  = RGBColor(0x9B, 0x59, 0xB6)
ACCENT_TEAL    = RGBColor(0x00, 0x8B, 0x8B)
LIGHT_BLUE     = RGBColor(0xDD, 0xEB, 0xF7)
LIGHT_GREEN    = RGBColor(0xE2, 0xEF, 0xDA)
LIGHT_ORANGE   = RGBColor(0xFC, 0xE4, 0xD6)
LIGHT_YELLOW   = RGBColor(0xFF, 0xF2, 0xCC)
LIGHT_PURPLE   = RGBColor(0xEF, 0xE8, 0xF9)
LIGHT_TEAL     = RGBColor(0xE0, 0xF7, 0xFA)
WHITE          = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK      = RGBColor(0x2D, 0x34, 0x36)   # dark grey (not black)
TEXT_MID       = RGBColor(0x63, 0x6E, 0x72)
SECTION_COLORS = {
    "Opening":       ACCENT_BLUE,
    "Foundations":   ACCENT_BLUE,
    "Tools":         ACCENT_TEAL,
    "Spec-Driven":   ACCENT_GREEN,
    "Architecture":  ACCENT_ORANGE,
    "Implementation":ACCENT_PURPLE,
    "Quality":       ACCENT_GREEN,
    "Security":      RGBColor(0xC0, 0x39, 0x2B),
    "DevOps":        ACCENT_ORANGE,
    "Performance":   ACCENT_TEAL,
    "Multi-Agent":   ACCENT_PURPLE,
    "Governance":    ACCENT_BLUE,
    "Closing":       ACCENT_BLUE,
    "Backup":        TEXT_MID,
    "Practical":     ACCENT_GREEN,
}

# ── Reusable helpers ───────────────────────────────────────────────────

def new_presentation():
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs

def set_bg(slide, color=BG_COLOR):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_slide(prs, bg_color=BG_COLOR):
    layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(layout)
    set_bg(slide, bg_color)
    return slide

def _set_text(tf, text, font_size, bold, color, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = ""
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    p.alignment = alignment

def add_textbox(slide, left, top, width, height, text, font_size=14,
                bold=False, color=TEXT_DARK, alignment=PP_ALIGN.LEFT,
                font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    _set_text(tf, text, font_size, bold, color, alignment, font_name)
    return txBox

def add_title_bar(slide, title_text, section="Opening", slide_num=1):
    """Blue bar across the top with title text."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), SLIDE_W, Inches(0.95)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_BAR
    bar.line.fill.background()
    tf = bar.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title_text
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = "Calibri"

    # Section pill bottom-left
    sec_color = SECTION_COLORS.get(section, ACCENT_BLUE)
    pill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.4), Inches(6.85), Inches(4.5), Inches(0.38)
    )
    pill.fill.solid()
    pill.fill.fore_color.rgb = sec_color
    pill.line.fill.background()
    ptf = pill.text_frame
    ptf.margin_left = Inches(0.15)
    ptf.margin_top = Inches(0.02)
    pr = ptf.paragraphs[0].add_run()
    pr.text = f"AI-Native Software Development  |  {section}"
    pr.font.size = Pt(10)
    pr.font.color.rgb = WHITE
    pr.font.name = "Calibri"

    # Slide number bottom-right
    add_textbox(slide, 12.4, 6.85, 0.7, 0.35, str(slide_num),
                font_size=11, bold=True, color=TEXT_MID,
                alignment=PP_ALIGN.RIGHT)

def add_bullet_list(slide, left, top, width, height, items,
                    font_size=16, color=TEXT_DARK, spacing=Pt(8),
                    bullet_char="\u2022", line_spacing=1.3):
    """Add a bulleted list. items can be str or (str, sub_items_list)."""
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                     Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        sub_items = []
        if isinstance(item, tuple):
            item, sub_items = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = spacing
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = f"{bullet_char} {item}"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        for si in sub_items:
            sp = tf.add_paragraph()
            sp.space_after = Pt(4)
            sp.line_spacing = line_spacing
            sr = sp.add_run()
            sr.text = f"     - {si}"
            sr.font.size = Pt(font_size - 2)
            sr.font.color.rgb = TEXT_MID
            sr.font.name = "Calibri"
    return txBox

def add_rounded_box(slide, left, top, width, height, text,
                    fill_color=LIGHT_BLUE, text_color=TEXT_DARK,
                    font_size=11, bold=False, border_color=None, alignment=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = alignment
    # Support multi-line with \n
    lines = text.split("\n")
    for i, line in enumerate(lines):
        if i > 0:
            p = tf.add_paragraph()
            p.alignment = alignment
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold if i == 0 else False
        run.font.color.rgb = text_color
        run.font.name = "Calibri"
    return shape

def add_arrow_right(slide, left, top, width=0.5, height=0.01, color=ACCENT_BLUE):
    """Horizontal arrow pointing right."""
    arr = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(left), Inches(top), Inches(width), Inches(0.3)
    )
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    return arr

def add_arrow_down(slide, left, top, width=0.3, height=0.4, color=ACCENT_BLUE):
    arr = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    arr.fill.solid()
    arr.fill.fore_color.rgb = color
    arr.line.fill.background()
    return arr

def add_connector_line(slide, x1, y1, x2, y2, color=ACCENT_BLUE, width=1.5):
    """Add a simple straight line connector."""
    conn = slide.shapes.add_connector(
        1,  # straight connector
        Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    return conn

def add_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

def add_code_box(slide, left, top, width, height, code_text,
                 font_size=10, bg_color=RGBColor(0xF8,0xF9,0xFA)):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = RGBColor(0xDD,0xDD,0xDD)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_top = Inches(0.1)
    tf.margin_right = Inches(0.1)
    lines = code_text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.name = "Consolas"
        run.font.color.rgb = TEXT_DARK
        p.space_after = Pt(1)
    return shape

def add_table_shape(slide, left, top, width, height, rows, cols,
                    header_color=TITLE_BAR, row_colors=None):
    """Add a table with data. rows is list of lists."""
    table_shape = slide.shapes.add_table(len(rows), cols,
                                          Inches(left), Inches(top),
                                          Inches(width), Inches(height))
    table = table_shape.table
    for r, row_data in enumerate(rows):
        for c, cell_text in enumerate(row_data):
            cell = table.cell(r, c)
            cell.text = str(cell_text)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(12)
                    run.font.name = "Calibri"
                    if r == 0:
                        run.font.bold = True
                        run.font.color.rgb = WHITE
                    else:
                        run.font.color.rgb = TEXT_DARK
            # Header row coloring
            if r == 0:
                _set_cell_fill(cell, header_color)
            else:
                if row_colors and r - 1 < len(row_colors):
                    _set_cell_fill(cell, row_colors[r - 1])
                else:
                    _set_cell_fill(cell, WHITE if r % 2 == 1 else LIGHT_BLUE)
    return table_shape

def _set_cell_fill(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    solidFill = tcPr.makeelement(qn('a:solidFill'), {})
    srgbClr = solidFill.makeelement(qn('a:srgbClr'), {'val': str(color)})
    solidFill.append(srgbClr)
    # Remove existing fills
    for child in list(tcPr):
        if child.tag.endswith('solidFill'):
            tcPr.remove(child)
    tcPr.append(solidFill)

def save_deck(prs, filename):
    path = os.path.join("/home/gurukraj/github/faculty-program", filename)
    prs.save(path)
    print(f"Saved: {path}")
    return path
